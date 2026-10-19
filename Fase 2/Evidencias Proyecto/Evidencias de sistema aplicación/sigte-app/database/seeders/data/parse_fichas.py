"""Lee el Excel del hospital y deja fichas con código, número y código de cada pieza."""

import json
import re
import unicodedata
from pathlib import Path

import openpyxl

EXCEL = Path(r"C:\Users\ymaureira\Downloads\INVENTARIO CAJAS EXCEL.xlsx")
SALIDA = Path(__file__).with_name("fichas.json")
ANTERIOR = Path(__file__).with_name("fichas.anterior.json")

PARAR = (
    "TOTAL",
    "RESPONSABLE",
    "CONTROL INTERNO",
    "PRECINTO",
    "OBSERVACION",
    "FIRMA",
)


def plano(valor) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().upper()


def limpio(valor) -> str:
    texto = "" if valor is None else str(valor)
    texto = re.sub(r"\s+", " ", texto).strip()
    texto = re.sub(r"^\*?\s*nombre\s*[:\-–—]+\s*caja\s*:\s*", "", texto, flags=re.I)
    texto = re.sub(r"^\*?\s*nombre\s+caja\s*:\s*", "", texto, flags=re.I)
    texto = re.sub(r"^\*?\s*nombre\s*[:\-–—]+\s*", "", texto, flags=re.I)
    texto = re.sub(r"^\*?\s*nombre\s+", "", texto, flags=re.I)
    return texto.strip(" -:")


def es_nombre(celda: str) -> bool:
    base = plano(celda).rstrip(" .")
    return base in {"INSTRUMENTAL", "DESCRIPCION", "PRODUCTO"} or base.startswith("INSTRUMENTAL ")


def es_codigo(celda: str) -> bool:
    return plano(celda).startswith("CODIGO")


def es_cantidad(celda: str) -> bool:
    base = plano(celda).rstrip(" .")
    return base in {"CANT", "CANTIDAD"}


def es_tamano(celda: str) -> bool:
    return plano(celda).startswith("TAMAN")


def parece_codigo(valor: str) -> bool:
    token = valor.strip()
    if not token or " " in token or len(token) > 40 or len(token) < 2:
        return False
    if not re.search(r"\d", token):
        return False
    return bool(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9./_-]*", token))


def numero_en(textos) -> str:
    for texto in textos:
        if not texto:
            continue
        hallado = re.search(r"(?:NUMERO|NRO|N\s*[°º])\s*:?\s*#?\s*0*(\d+)", plano(texto))
        if hallado:
            return hallado.group(1)
        hallado = re.search(r"#\s*0*(\d+)", str(texto))
        if hallado:
            return hallado.group(1)
        hallado = re.search(r"\bN\s*[°º]\s*0*(\d+)", plano(texto))
        if hallado:
            return hallado.group(1)
    return ""


def codigo_caja(textos) -> str:
    for texto in textos:
        if not texto:
            continue
        hallado = re.search(r"CODIGO\s+([A-Za-z0-9][A-Za-z0-9./_-]*)", plano(texto))
        if hallado and not hallado.group(1).isdigit():
            return hallado.group(1)
    return ""


def separar_codigo_del_nombre(nombre: str) -> tuple[str, str]:
    partes = nombre.split(" ", 1)
    if len(partes) == 2 and parece_codigo(partes[0]) and not parece_codigo(partes[1][:1] or "x"):
        return partes[1].strip(), partes[0]
    return nombre, ""


def hoja_filas(ws, max_col=12):
    filas = []
    for row in ws.iter_rows(max_col=max_col, values_only=True):
        filas.append(list(row))
    while filas and all(c in (None, "") for c in filas[-1]):
        filas.pop()
    return filas


def bloques(filas):
    encontrados = []
    for i, fila in enumerate(filas):
        nombres = [c for c, valor in enumerate(fila) if es_nombre(str(valor or ""))]
        if not nombres:
            continue
        limites = nombres + [len(fila)]
        for inicio, fin in zip(nombres, limites[1:]):
            if fin <= inicio:
                continue
            codigo = cantidad = tamano = None
            for col in range(inicio, fin):
                valor = str(fila[col] or "")
                if es_codigo(valor) and codigo is None:
                    codigo = col
                elif es_cantidad(valor) and cantidad is None:
                    cantidad = col
                elif es_tamano(valor) and tamano is None:
                    tamano = col
            if cantidad is None:
                for salto in (1, 2):
                    if i + salto >= len(filas):
                        break
                    for col in range(max(0, inicio - 1), fin + 1):
                        if col < len(filas[i + salto]) and es_cantidad(str(filas[i + salto][col] or "")):
                            cantidad = col
                            break
                    if cantidad is not None:
                        break
            if codigo is None:
                for col in range(max(0, inicio - 2), inicio):
                    if es_codigo(str(fila[col] or "")):
                        codigo = col
            encontrados.append(
                {
                    "fila": i,
                    "nombre": inicio,
                    "fin": fin,
                    "codigo": codigo,
                    "cantidad": cantidad,
                    "tamano": tamano,
                }
            )
    return encontrados


def es_ruido(texto: str) -> bool:
    clave = plano(texto)
    if not clave or es_nombre(texto) or es_cantidad(texto) or es_codigo(texto) or clave in PARAR:
        return True
    if clave.startswith("NUMERO") or clave.startswith("NUMERO:") or clave.startswith("N "):
        return True
    if clave.startswith("PEDIDO") or clave.startswith("GUIA DE DESPACHO") or clave.startswith("FUNCIONARIO") or clave.startswith("RESPONSABLE"):
        return True
    return clave in {"FECHA", "L.", "+", "OBS"}


def titulo_de(filas, bloque):
    candidatos = []
    for i in range(bloque["fila"] - 1, max(-1, bloque["fila"] - 8), -1):
        fila = filas[i]
        for col in range(max(0, bloque["nombre"] - 2), min(len(fila), bloque["fin"] + 1)):
            valor = fila[col]
            if valor in (None, ""):
                continue
            texto = str(valor).strip()
            if es_ruido(texto):
                continue
            candidatos.append(texto)
    nombre = ""
    for texto in candidatos:
        if any(marca in plano(texto) for marca in ("NOMBRE", "CAJA", "SET")):
            nombre = limpio(texto)
            break
    if nombre == "" and candidatos:
        nombre = limpio(candidatos[0])
    alrededor = []
    for i in range(max(0, bloque["fila"] - 6), bloque["fila"] + 1):
        alrededor.extend(str(c).strip() for c in filas[i] if c not in (None, ""))
    return nombre, numero_en(alrededor), codigo_caja(alrededor)


def items_de(filas, bloque):
    items = []
    vacios = 0
    for i in range(bloque["fila"] + 1, len(filas)):
        fila = filas[i]
        if any(es_nombre(str(c or "")) for c in fila):
            break
        nombre = fila[bloque["nombre"]] if bloque["nombre"] < len(fila) else None
        if nombre in (None, ""):
            vacios += 1
            if items and vacios >= 3:
                break
            continue
        vacios = 0
        texto = str(nombre).strip()
        clave = plano(texto)
        if clave in PARAR or clave.startswith("RESPONSABLE") or es_cantidad(texto) or es_codigo(texto):
            if items:
                break
            continue
        cantidad = 1
        if bloque["cantidad"] is not None and bloque["cantidad"] < len(fila) and fila[bloque["cantidad"]] not in (None, ""):
            crudo = fila[bloque["cantidad"]]
            if isinstance(crudo, (int, float)) and not isinstance(crudo, bool):
                cantidad = int(crudo)
            else:
                hallado = re.search(r"\d+", str(crudo))
                cantidad = int(hallado.group(0)) if hallado else 0
        if cantidad <= 0:
            continue
        codigo = ""
        if bloque["codigo"] is not None and bloque["codigo"] < len(fila) and fila[bloque["codigo"]] not in (None, ""):
            codigo = str(fila[bloque["codigo"]]).strip()
        tamano = ""
        if bloque["tamano"] is not None and bloque["tamano"] < len(fila) and fila[bloque["tamano"]] not in (None, ""):
            tamano = str(fila[bloque["tamano"]]).strip()
        if codigo == "":
            texto, codigo = separar_codigo_del_nombre(texto)
        if texto == "":
            continue
        items.append(
            {
                "nombre": texto[:160],
                "tamano": tamano[:40],
                "codigo": codigo[:80],
                "cantidad": max(1, cantidad),
            }
        )
    return items


def sin_encabezado(filas, hoja):
    """Hojas que traen nombre, código y cantidad sin fila de títulos."""
    fichas = []
    actual = None
    for fila in filas:
        celdas = [("" if c is None else str(c).strip()) for c in fila]
        if not any(celdas):
            continue
        if len(celdas) >= 3 and parece_codigo(celdas[1]) and re.fullmatch(r"\d+", celdas[2] or ""):
            if actual is None:
                actual = {"hoja": hoja, "nombre": hoja, "numero": "", "codigo_referencia": "", "tipo": "caja", "items": []}
            actual["items"].append(
                {"nombre": celdas[0][:160], "tamano": "", "codigo": celdas[1][:80], "cantidad": int(celdas[2])}
            )
            continue
        if actual and actual["items"]:
            fichas.append(actual)
            actual = None
        if any(celdas) and not parece_codigo(celdas[0]):
            referencia = celdas[1] if len(celdas) > 1 and parece_codigo(celdas[1]) else ""
            actual = {
                "hoja": hoja,
                "nombre": limpio(celdas[0]) or hoja,
                "numero": numero_en(celdas),
                "codigo_referencia": referencia,
                "tipo": "set" if "SET" in plano(celdas[0]) else "caja",
                "items": [],
            }
    if actual and actual["items"]:
        fichas.append(actual)
    return fichas


def parsear():
    libro = openpyxl.load_workbook(EXCEL, read_only=True, data_only=True)
    fichas = []
    for nombre_hoja in libro.sheetnames:
        hoja = re.sub(r"\s+", " ", nombre_hoja).strip()
        filas = hoja_filas(libro[nombre_hoja])
        vistos = set()
        grupos = bloques(filas)
        for bloque in grupos:
            nombre, numero, referencia = titulo_de(filas, bloque)
            items = items_de(filas, bloque)
            if not items:
                continue
            if nombre == "" or es_ruido(nombre):
                nombre = hoja
            if es_ruido(nombre):
                continue
            firma = (plano(nombre), numero, tuple(plano(i["nombre"]) for i in items[:3]))
            if firma in vistos:
                continue
            vistos.add(firma)
            fichas.append(
                {
                    "hoja": hoja,
                    "nombre": nombre[:160],
                    "numero": numero,
                    "codigo_referencia": referencia,
                    "tipo": "set" if re.search(r"\bSET\b", plano(nombre)) else "caja",
                    "items": items,
                }
            )
        if not any(f["hoja"] == hoja for f in fichas):
            fichas.extend(sin_encabezado(filas, hoja))
    return fichas


def asignar_codigos(fichas):
    anteriores = json.loads(ANTERIOR.read_text(encoding="utf-8")) if ANTERIOR.exists() else []
    por_clave = {}
    for ficha in anteriores:
        clave = (plano(ficha.get("hoja", "")), plano(limpio(ficha.get("nombre", ""))))
        por_clave.setdefault(clave, []).append(ficha.get("codigo"))
    usados = set()
    mayor = 0
    for codigo in (f.get("codigo") or "" for f in anteriores):
        hallado = re.fullmatch(r"FIC-(\d+)", codigo or "")
        if hallado:
            mayor = max(mayor, int(hallado.group(1)))
            usados.add(codigo)

    def siguiente():
        nonlocal mayor
        while True:
            mayor += 1
            codigo = f"FIC-{mayor:03d}"
            if codigo not in usados:
                usados.add(codigo)
                return codigo

    for ficha in fichas:
        clave = (plano(ficha["hoja"]), plano(ficha["nombre"]))
        bolsa = por_clave.get(clave) or []
        ficha["codigo"] = bolsa.pop(0) if bolsa else siguiente()

    piezas = {}
    correlativo = 1
    for ficha in fichas:
        for item in ficha["items"]:
            if item["codigo"]:
                continue
            clave = (plano(item["nombre"]), plano(item["tamano"]))
            if clave not in piezas:
                piezas[clave] = f"INS-{correlativo:05d}"
                correlativo += 1
            item["codigo"] = piezas[clave]
    return fichas


def main():
    fichas = asignar_codigos(parsear())
    SALIDA.write_text(json.dumps(fichas, ensure_ascii=False), encoding="utf-8")
    con_ref = sum(1 for f in fichas if f["codigo_referencia"])
    con_num = sum(1 for f in fichas if f["numero"])
    lineas = sum(len(f["items"]) for f in fichas)
    excel = sum(1 for f in fichas for i in f["items"] if not i["codigo"].startswith("INS-"))
    print(f"fichas={len(fichas)} lineas={lineas} con_codigo_excel={excel} con_numero={con_num} con_referencia={con_ref}")


if __name__ == "__main__":
    main()
